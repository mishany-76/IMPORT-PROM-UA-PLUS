import os
import io
import gzip
import boto3
import requests
import json
import time
from openpyxl import load_workbook
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

# --- НАСТРОЙКИ ДЛЯ ЭШЕЛОНА 2 (ТЯЖЕЛАЯ ТАБЛИЦА) ---
# Подгружаем данные из переменных окружения
R2_ACCESS_KEY = os.getenv('R2_ACCESS_KEY', 'R2_ACCESS_KEY')
R2_SECRET_KEY = os.getenv('R2_SECRET_KEY', 'R2_SECRET_KEY')
R2_ACCOUNT_ID = os.getenv('R2_ACCOUNT_ID', 'R2_ACCOUNT_ID')
R2_BUCKET_NAME = os.getenv('R2_BUCKET_NAME', 'R2_BUCKET_NAME')
R2_ENDPOINT = f'https://{R2_ACCOUNT_ID}.r2.cloudflarestorage.com'
R2_PUBLIC_DOMAIN = os.getenv('R2_PUBLIC_DOMAIN', 'R2_PUBLIC_DOMAIN')

DELAY_BEFORE_FIRST = int(os.getenv('DELAY_BEFORE_FIRST', 10))
DELAY_BETWEEN_REQUESTS = int(os.getenv('DELAY_BETWEEN_REQUESTS', 20))

# Имя файла на Cloudflare R2 для тяжелого каталога магазина VaillantVibe
FILE_NAME = 'import_data_heavy_vaillant.xlsx.gz'
PUBLIC_URL = f'{R2_PUBLIC_DOMAIN}/{FILE_NAME}'

# Конфигурация магазинов кластера Prom.ua (CityStyle, BestSeller, VaillantVibe)
# Токены подтягиваются из переменных окружения / GitHub Secrets
PROM_STORES = [
    {
        "name": "CityStyle (Флагман)",
        "token": os.getenv('PROM_TOKEN_City') or os.getenv('PROM_TOKEN_CITY') or ''
    },
    {
        "name": "BestSeller (Сателлит)",
        "token": os.getenv('PROM_TOKEN_Best') or os.getenv('PROM_TOKEN_BEST') or ''
    },
    {
        "name": "VaillantVibe (Сателлит)",
        "token": os.getenv('PROM_TOKEN_Vaillant') or os.getenv('PROM_TOKEN_VAILLANT') or ''
    },
]

# ID Тяжелой таблицы-контент-банка (Эшелон 2)
SPREADSHEET_ID = os.getenv('IMPORT_PROM_UA_HEAVY_SPREADSHEET_ID', '1Rkbne3OmHOC3vZtvnuvW_eINL1dBIKdNI7sYkMxATmM')
SERVICE_ACCOUNT_FILE = 'key_sheet.json'


def get_gspread_xlsx():
    """Скачивает тяжелую Google таблицу в формате XLSX в память"""
    print(f"1. Авторизация в Google и скачивание XLSX (Таблица ID: {SPREADSHEET_ID})...")
    creds = Credentials.from_service_account_file(
        SERVICE_ACCOUNT_FILE,
        scopes=['https://www.googleapis.com/auth/drive.readonly']
    )
    service = build('drive', 'v3', credentials=creds)

    request = service.files().export_media(
        fileId=SPREADSHEET_ID,
        mimeType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    done = False
    while done is False:
        status, done = downloader.next_chunk()
    return fh.getvalue()


def process_xlsx(xlsx_content):
    """Удаляет лист 'Systemic' и сжимает в Gzip"""
    print("2. Обработка файла: удаление служебного листа 'Systemic'...")
    wb = load_workbook(io.BytesIO(xlsx_content))

    if "Systemic" in wb.sheetnames:
        del wb["Systemic"]
        print("   Лист 'Systemic' удален.")

    # Сохраняем обработанный XLSX в буфер
    processed_xlsx_bin = io.BytesIO()
    wb.save(processed_xlsx_bin)

    print("3. Сжатие файла в Gzip...")
    gzip_buffer = io.BytesIO()
    with gzip.GzipFile(fileobj=gzip_buffer, mode='wb') as f:
        f.write(processed_xlsx_bin.getvalue())

    return gzip_buffer.getvalue()


def upload_to_r2(data):
    """Загрузка файла в Cloudflare R2"""
    print(f"4. Загрузка в Cloudflare R2 (Бакет: {R2_BUCKET_NAME}, Файл: {FILE_NAME})...")
    s3 = boto3.client(
        's3',
        endpoint_url=R2_ENDPOINT,
        aws_access_key_id=R2_ACCESS_KEY,
        aws_secret_access_key=R2_SECRET_KEY,
        region_name='auto'
    )

    s3.put_object(
        Bucket=R2_BUCKET_NAME,
        Key=FILE_NAME,
        Body=data,
        ContentType='application/x-gzip',
        ACL='public-read'
    )
    print(f"   Файл доступен по ссылке: {PUBLIC_URL}")


def trigger_prom_import():
    """
    Запуск асинхронного импорта контента в магазины Prom.ua по строго утвержденному стандарту Эшелона 2.
    Отправляет POST https://my.prom.ua/api/v1/products/import_url со ссылкой на архив на Cloudflare R2.
    """
    import_url = "https://my.prom.ua/api/v1/products/import_url"

    # СТАНДАРТ ЭШЕЛОНА 2:
    # only_update = True  -> обновлять только существующие позиции (не плодить дубли)
    # mark_missing_product_as = 'none' -> отсутствующие товары НЕ удалять и НЕ скрывать
    # force_update = True -> принудительное обновление контента
    # updated_fields -> строгий белый список (SEO-слова, названия, описания с таблицами)
    payload = {
        "url": PUBLIC_URL,
        "force_update": True,
        "only_available": False,
        "only_update": True,
        "mark_missing_product_as": "none",
        "updated_fields": [
            "keywords",
            "keywords_ua",
            "name",
            "name_ua",
            "description",
            "description_ua"
        ]
    }

    print("5. Автоматическая отправка ссылки на архив в Prom API (Эшелон 2)...")
    print(f"   Публичный URL архива: {PUBLIC_URL}")

    is_first = True
    active_dispatches = 0

    for store in PROM_STORES:
        store_name = store["name"]
        token = (store.get("token") or "").strip()

        # Проверка валидности токена
        if not token or token.startswith("PROM_TOKEN") or token in ["токен_1", ""]:
            print(f"   [Пропуск] {store_name}: API-токен не настроен в GitHub Secrets.")
            continue

        if is_first:
            print(f"   Ожидание {DELAY_BEFORE_FIRST} сек. перед первой отправкой...")
            time.sleep(DELAY_BEFORE_FIRST)
            is_first = False
        else:
            print(f"   Ожидание {DELAY_BETWEEN_REQUESTS} сек. перед следующим запросом...")
            time.sleep(DELAY_BETWEEN_REQUESTS)

        headers = {
            'Authorization': f'Bearer {token}',
            'Content-Type': 'application/json',
            'X-Language': 'uk'
        }

        try:
            print(f"   -> Отправка запроса на импорт для {store_name}...")
            response = requests.post(import_url, json=payload, headers=headers, timeout=60)
            if response.status_code == 200:
                data = response.json()
                import_id = data.get('id', 'N/A')
                print(f"   [OK] {store_name}: Импорт успешно запущен на Prom.ua! ID задачи: {import_id}")
                active_dispatches += 1
            else:
                print(f"   [Error] {store_name}: HTTP {response.status_code} - {response.text}")
        except Exception as e:
            print(f"   [Exception] Ошибка при запросе к Prom.ua для {store_name}: {e}")

    if active_dispatches == 0:
        print("   [Инфо] Ни один активный токен магазина не был обработан (проверьте секреты PROM_TOKEN_*).")
    else:
        print(f"   [Успех] Запущено задач импорта: {active_dispatches}.")


if __name__ == "__main__":
    try:
        raw_xlsx = get_gspread_xlsx()
        gzipped_xlsx = process_xlsx(raw_xlsx)
        upload_to_r2(gzipped_xlsx)
        trigger_prom_import()
        print("\nГотово! Тяжелый архив сформирован, выгружен на R2 и отправлен в Prom API.")
    except Exception as e:
        print(f"\nКритическая ошибка: {e}")
